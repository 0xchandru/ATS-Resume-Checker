import logging
import os
from typing import Any

logger = logging.getLogger(__name__)


def parse_resume(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext == ".docx":
        return _parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _parse_pdf(file_path: str) -> dict:
    try:
        import pdfplumber
        return _parse_pdf_pdfplumber(file_path)
    except Exception as e:
        logger.warning("pdfplumber failed (%s), falling back to PyMuPDF", e)
        return _parse_pdf_pymupdf(file_path)


def _parse_pdf_pdfplumber(file_path: str) -> dict:
    import pdfplumber

    pages_data = []
    all_text_parts = []
    char_data = []
    fonts_seen = set()
    has_tables = False
    has_header_footer = False
    has_images = False

    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page_num, page in enumerate(pdf.pages):
            page_height = page.height or 1000

            text = page.extract_text() or ""
            all_text_parts.append(text)

            tables = page.extract_tables()
            if tables and len(tables) > 0 and any(t for t in tables):
                has_tables = True

            chars = page.chars
            for ch in chars:
                fn = ch.get("fontname", "")
                fs = ch.get("size", 0)
                if fn:
                    fonts_seen.add((fn, round(fs, 1)))
                top = ch.get("top", 0)
                bottom = ch.get("bottom", 0)
                if top < page_height * 0.08 or bottom > page_height * 0.92:
                    has_header_footer = True
                char_data.append({"x": ch.get("x0", 0), "y": ch.get("top", 0), "fontname": fn, "size": fs, "text": ch.get("text", "")})

            pages_data.append({"page": page_num + 1, "text": text, "char_count": len(chars)})

    raw_text = "\n".join(all_text_parts).strip()
    fonts_list = [{"fontname": fn, "size": fs} for fn, fs in fonts_seen]

    try:
        import fitz
        doc = fitz.open(file_path)
        for page in doc:
            imgs = page.get_images(full=True)
            if imgs:
                has_images = True
                break
        doc.close()
    except Exception:
        pass

    return {
        "raw_text": raw_text,
        "pages": pages_data,
        "page_count": page_count,
        "file_type": "pdf",
        "metadata": {
            "fonts_used": [f"{f['fontname']}" for f in fonts_list],
            "font_count": len(fonts_seen),
            "font_details": fonts_list,
            "has_tables": has_tables,
            "has_images": has_images,
            "has_header_footer": has_header_footer,
            "has_special_characters": _has_special_chars(raw_text),
            "char_data": char_data[:500],
            "word_count": len(raw_text.split()),
        },
    }


def _parse_pdf_pymupdf(file_path: str) -> dict:
    import fitz

    pages_data = []
    all_text_parts = []
    has_images = False

    doc = fitz.open(file_path)
    page_count = len(doc)

    for page_num, page in enumerate(doc):
        text = page.get_text("text")
        all_text_parts.append(text)
        pages_data.append({"page": page_num + 1, "text": text, "char_count": len(text)})
        if page.get_images(full=True):
            has_images = True

    doc.close()
    raw_text = "\n".join(all_text_parts).strip()

    return {
        "raw_text": raw_text,
        "pages": pages_data,
        "page_count": page_count,
        "file_type": "pdf",
        "metadata": {
            "fonts_used": [],
            "font_count": 0,
            "font_details": [],
            "has_tables": False,
            "has_images": has_images,
            "has_header_footer": False,
            "has_special_characters": _has_special_chars(raw_text),
            "char_data": [],
            "word_count": len(raw_text.split()),
        },
    }


def _parse_docx(file_path: str) -> dict:
    from docx import Document
    import lxml.etree as etree

    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    raw_text = "\n".join(paragraphs)

    styles = set()
    for para in doc.paragraphs:
        if para.style and para.style.name:
            styles.add(para.style.name)

    has_header_footer = False
    try:
        for section in doc.sections:
            header_text = " ".join(p.text for p in section.header.paragraphs).strip()
            footer_text = " ".join(p.text for p in section.footer.paragraphs).strip()
            if header_text or footer_text:
                has_header_footer = True
                break
    except Exception:
        pass

    has_tables = len(doc.tables) > 0
    has_images = len(doc.inline_shapes) > 0

    has_text_boxes = False
    try:
        body_xml = doc.element.body.xml
        has_text_boxes = "txbxContent" in body_xml
    except Exception:
        pass

    return {
        "raw_text": raw_text,
        "pages": [{"page": 1, "text": raw_text, "char_count": len(raw_text)}],
        "page_count": 1,
        "file_type": "docx",
        "metadata": {
            "fonts_used": list(styles),
            "font_count": len(styles),
            "font_details": [],
            "has_tables": has_tables or has_text_boxes,
            "has_images": has_images,
            "has_header_footer": has_header_footer,
            "has_special_characters": _has_special_chars(raw_text),
            "char_data": [],
            "word_count": len(raw_text.split()),
            "has_text_boxes": has_text_boxes,
        },
    }


def _has_special_chars(text: str) -> bool:
    # Only flag severe PDF extraction failures, not standard typographic symbols like bullets/dashes
    for ch in text:
        if ch == '' or (ord(ch) < 32 and ch not in '\n\r\t'):
            return True
    return False
